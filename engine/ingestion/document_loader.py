from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import List


KB_ROOT = Path(
    "knowledge_base/ANAHAT_KnowledgeBase_v3"
)


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
}


@dataclass
class LoadedDocument:
    document_id: str
    source_path: str
    source_file: str
    extension: str
    text: str


def make_document_id(path: Path) -> str:
    """
    Create a stable document ID from the path relative
    to the ANAHAT knowledge-base root.

    Example:

        README.md
            -> README

        raaga/README.md
            -> raaga_README

        rag/chakra_reasoning/root_chakra_reasoning.md
            -> rag_chakra_reasoning_root_chakra_reasoning
    """

    try:
        relative = path.relative_to(KB_ROOT)
    except ValueError:
        relative = path

    relative_without_suffix = relative.with_suffix("")

    document_id = "_".join(
        relative_without_suffix.parts
    )

    document_id = re.sub(
        r"[^A-Za-z0-9_.-]+",
        "_",
        document_id,
    )

    return document_id


def load_text_file(path: Path) -> str:

    return path.read_text(
        encoding="utf-8",
        errors="replace",
    )


def load_pdf(path: Path) -> str:

    from pypdf import PdfReader

    reader = PdfReader(str(path))

    pages = []

    for page in reader.pages:

        text = page.extract_text() or ""

        pages.append(text)

    return "\n\n".join(pages)


def load_docx(path: Path) -> str:

    from docx import Document

    document = Document(str(path))

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    return "\n\n".join(paragraphs)


def load_document(path: Path) -> LoadedDocument:

    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:

        raise ValueError(
            f"Unsupported document type: {suffix}"
        )

    if suffix in {".txt", ".md"}:

        text = load_text_file(path)

    elif suffix == ".pdf":

        text = load_pdf(path)

    elif suffix == ".docx":

        text = load_docx(path)

    else:

        raise ValueError(
            f"No loader available for {suffix}"
        )

    document_id = make_document_id(path)

    return LoadedDocument(
        document_id=document_id,
        source_path=str(path),
        source_file=path.name,
        extension=suffix,
        text=text,
    )


def discover_documents(root: str) -> List[Path]:

    root_path = Path(root)

    documents = []

    for path in root_path.rglob("*"):

        if not path.is_file():
            continue

        if path.suffix.lower() in SUPPORTED_EXTENSIONS:

            documents.append(path)

    return sorted(documents)